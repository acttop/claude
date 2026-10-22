"""폴더 일괄 변환 테스트: 입력 폴더의 모든 PDF를 변환하고 결과 요약 리포트를 만든다.

사용: python batch_test.py [입력폴더] [출력폴더]
  기본값: 입력 D:\\pdftoexcel, 출력 D:\\pdftoexcel\\results (작업폴더)
"""
from __future__ import annotations

import sys
import time
import traceback
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from pdf2excel.core import convert
from pdf2excel.core import _autofit  # 리포트 열 너비 맞춤 재사용

DEFAULT_IN = Path(r"D:\pdftoexcel")
DEFAULT_OUT = DEFAULT_IN / "results"


def main() -> int:
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_IN
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUT
    pdfs = sorted(p for p in src.rglob("*") if p.suffix.lower() == ".pdf"
                  and out.resolve() not in p.resolve().parents)  # 결과 폴더 재처리 방지
    if not pdfs:
        print(f"[오류] PDF가 없습니다: {src}")
        return 1
    out.mkdir(parents=True, exist_ok=True)
    print(f"입력: {src} ({len(pdfs)}개)\n출력: {out}\n")

    rows = []
    for pdf in pdfs:
        target = out / pdf.relative_to(src).with_suffix(".xlsx")
        t0 = time.perf_counter()
        try:
            r = convert(pdf, target)
            status = "주의" if r.warnings else "성공"
            methods = ", ".join(f"{p}쪽:{m}" for p, m in r.methods.items())
            row = [pdf.name, status, r.pages, r.tables, r.text_pages, r.ocr_pages,
                   ",".join(map(str, r.failed_pages)), methods,
                   "; ".join(r.warnings), str(r.output)]
        except Exception as e:  # noqa: BLE001
            status = "실패"
            row = [pdf.name, status, None, None, None, None, "", "",
                   f"{type(e).__name__}: {e}", ""]
            (out / "errors.log").open("a", encoding="utf-8").write(
                f"== {pdf}\n{traceback.format_exc()}\n")
        row.insert(9, round(time.perf_counter() - t0, 2))
        rows.append(row)
        print(f"[{status}] {pdf.name}  페이지 {row[2]}, 표 {row[3]}, OCR {row[5]}, {row[9]}초")

    wb = Workbook()
    ws = wb.active
    ws.title = "테스트결과"
    ws.append(["파일명", "결과", "페이지", "표 개수", "텍스트 페이지", "OCR 페이지",
               "OCR 실패 페이지", "페이지별 방법", "경고/오류", "소요(초)", "출력 파일"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in rows:
        ws.append(row)
    _autofit(ws)
    report = out / f"테스트리포트_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    wb.save(report)

    cnt = {s: sum(r[1] == s for r in rows) for s in ("성공", "주의", "실패")}
    print(f"\n성공 {cnt['성공']} / 주의 {cnt['주의']} / 실패 {cnt['실패']}  → 리포트: {report}")
    return 1 if cnt["실패"] else 0


if __name__ == "__main__":
    sys.exit(main())
