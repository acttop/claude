"""PDF -> Excel 변환 핵심 로직 (외부 네트워크 접근 없음)."""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .analyze import METHOD_FALLBACK, METHOD_OCR, METHOD_OCR_FAILED, OCR_AUTO, OCR_OFF, analyze_page

log = logging.getLogger(__name__)

# 엑셀 셀에 허용되지 않는 제어문자 (openpyxl IllegalCharacterError 방지)
_ILLEGAL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_NUMBER = re.compile(r"^-?\(?[\d,]+(\.\d+)?\)?%?$")

MODE_TABLE = "table"  # 표만 추출
MODE_TEXT = "text"    # 페이지 텍스트를 줄 단위로 추출
MODE_AUTO = "auto"    # 표가 있으면 표, 없으면 텍스트

LAYOUT_PER_PAGE = "per_page"  # 페이지(표)마다 시트 생성
LAYOUT_SINGLE = "single"      # 하나의 시트에 이어 붙이기


@dataclass
class ConvertOptions:
    mode: str = MODE_AUTO
    layout: str = LAYOUT_PER_PAGE
    pages: list[int] | None = None  # 1부터 시작, None이면 전체
    convert_numbers: bool = True    # "1,234" -> 1234 숫자 변환
    password: str | None = None
    ocr: str = OCR_AUTO             # off / auto(텍스트 없는 페이지만) / force


@dataclass
class ConvertResult:
    output: Path
    pages: int = 0
    tables: int = 0
    text_pages: int = 0
    ocr_pages: int = 0
    failed_pages: list[int] = field(default_factory=list)  # method=ocr_failed 페이지
    methods: dict[int, str] = field(default_factory=dict)  # 페이지 → 추출 방법
    warnings: list[str] = field(default_factory=list)


def parse_pages(spec: str, max_page: int | None = None) -> list[int] | None:
    """'1-3,5' -> [1, 2, 3, 5]. 빈 문자열이면 None(전체)."""
    spec = (spec or "").strip()
    if not spec:
        return None
    pages: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            start, end = int(a), int(b) if b.strip() else (max_page or int(a))
            pages.update(range(start, end + 1))
        else:
            pages.add(int(part))
    if any(p < 1 for p in pages):
        raise ValueError("페이지 번호는 1 이상이어야 합니다.")
    return sorted(pages)


def _clean(value, convert_numbers: bool):
    if value is None:
        return None
    text = _ILLEGAL_CHARS.sub("", str(value)).strip()
    if not text:
        return None
    if convert_numbers and _NUMBER.match(text):
        raw = text.replace(",", "")
        negative = raw.startswith("(") and raw.endswith(")")
        raw = raw.strip("()")
        percent = raw.endswith("%")
        raw = raw.rstrip("%")
        try:
            num = float(raw)
        except ValueError:
            return text
        if negative:
            num = -num
        if percent:
            return num / 100
        return int(num) if num.is_integer() and "." not in raw else num
    return text


def _write_rows(ws, rows: Iterable[list], start_row: int, convert_numbers: bool,
                header: bool) -> int:
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    r = start_row
    for i, row in enumerate(rows):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=_clean(value, convert_numbers))
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, float) and str(value).strip().endswith("%"):
                cell.number_format = "0.##%"
            if header and i == 0:
                cell.font = Font(bold=True)
                cell.fill = head_fill
        r += 1
    return r


def _autofit(ws):
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            # 한글 등 전각 문자는 폭 2로 계산
            longest = max(
                sum(2 if ord(ch) > 0x2E80 else 1 for ch in line)
                for line in str(cell.value).splitlines() or [""]
            )
            widths[cell.column] = max(widths.get(cell.column, 0), longest)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(w + 2, 8), 60)


def _safe_sheet_title(title: str, used: set[str]) -> str:
    title = re.sub(r"[\\/*?:\[\]]", "_", title)[:31]
    base, n = title, 2
    while title in used:
        suffix = f"_{n}"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title


def _needs_ocr(page, ocr: str) -> bool:
    from .analyze import OCR_FORCE, has_text_layer
    return ocr == OCR_FORCE or (ocr != OCR_OFF and not has_text_layer(page))


def _open_engine():
    """OCR 엔진을 연다. 실패하면 (None, 사유)를 돌려 페이지 결과에 사유가 남게 한다."""
    try:
        from .extract_ocr import OcrEngine
        return OcrEngine(), None
    except Exception as e:  # noqa: BLE001
        return None, str(e)


def convert(pdf_path: str | Path, xlsx_path: str | Path | None = None,
            options: ConvertOptions | None = None,
            progress: Callable[[int, int], None] | None = None) -> ConvertResult:
    """PDF를 엑셀로 변환한다. progress(현재페이지, 전체페이지) 콜백 지원."""
    opts = options or ConvertOptions()
    pdf_path = Path(pdf_path)
    out = Path(xlsx_path) if xlsx_path else pdf_path.with_suffix(".xlsx")
    result = ConvertResult(output=out)

    wb = Workbook()
    wb.remove(wb.active)
    used_titles: set[str] = set()
    single_ws = None
    single_row = 1

    with pdfplumber.open(pdf_path, password=opts.password or "") as pdf:
        total = len(pdf.pages)
        targets = opts.pages or list(range(1, total + 1))
        invalid = [p for p in targets if p > total]
        if invalid:
            result.warnings.append(f"존재하지 않는 페이지 무시: {invalid} (전체 {total}쪽)")
        targets = [p for p in targets if p <= total]

        engine, engine_error = None, None
        for idx, page_no in enumerate(targets, start=1):
            page = pdf.pages[page_no - 1]
            if engine is None and engine_error is None and _needs_ocr(page, opts.ocr):
                engine, engine_error = _open_engine()

            pr = analyze_page(page, page_no, pdf_path=pdf_path, password=opts.password,
                              ocr=opts.ocr, engine=engine, engine_error=engine_error)
            result.methods[page_no] = pr.method
            if pr.method in (METHOD_OCR, METHOD_FALLBACK, METHOD_OCR_FAILED):
                result.ocr_pages += 1
            if pr.method == METHOD_FALLBACK:
                result.warnings.append(f"{page_no}쪽: {pr.error} → PDF 텍스트로 대체 추출")
            if pr.method == METHOD_OCR and pr.error:
                result.warnings.append(f"{page_no}쪽: {pr.error}")
            if pr.method == METHOD_OCR and pr.rotation:
                result.warnings.append(f"{page_no}쪽: {pr.rotation}도 회전된 스캔으로 판단해 보정")

            tables = pr.tables if opts.mode in (MODE_TABLE, MODE_AUTO) else []
            blocks: list[tuple[str, list[list], bool]] = []  # (이름, 행들, 헤더여부)
            if pr.method == METHOD_OCR_FAILED:
                # 조용한 실패 금지: 실패 사유를 결과와 시트에 남긴다
                result.failed_pages.append(page_no)
                result.warnings.append(f"{page_no}쪽: [ocr_failed] {pr.error}")
                rows = [[f"[OCR 실패] {pr.error}"]] + [[ln] for ln in pr.lines]
                blocks.append((f"P{page_no}_OCR실패", rows, False))
            elif tables:
                for t_i, table in enumerate(tables, start=1):
                    blocks.append((f"P{page_no}_표{t_i}", table, True))
                result.tables += len(tables)
            elif opts.mode in (MODE_TEXT, MODE_AUTO):
                if pr.lines:
                    blocks.append((f"P{page_no}_텍스트", [[ln] for ln in pr.lines], False))
                    result.text_pages += 1
                else:
                    result.warnings.append(f"{page_no}쪽: 추출할 텍스트가 없습니다")

            for name, rows, header in blocks:
                if opts.layout == LAYOUT_SINGLE:
                    if single_ws is None:
                        single_ws = wb.create_sheet(_safe_sheet_title("전체", used_titles))
                    label = single_ws.cell(row=single_row, column=1, value=f"[{name}]")
                    label.font = Font(bold=True, color="1F4E78")
                    single_row = _write_rows(single_ws, rows, single_row + 1,
                                             opts.convert_numbers, header) + 1
                else:
                    ws = wb.create_sheet(_safe_sheet_title(name, used_titles))
                    _write_rows(ws, rows, 1, opts.convert_numbers, header)
                    if header:
                        ws.freeze_panes = "A2"

            result.pages += 1
            if progress:
                progress(idx, len(targets))
        if engine is not None:
            engine.close()

    if not wb.sheetnames:
        ws = wb.create_sheet("결과없음")
        ws["A1"] = "추출된 데이터가 없습니다."
        result.warnings.append("추출된 데이터가 없습니다.")
    for ws in wb.worksheets:
        _autofit(ws)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    log.info("저장 완료: %s", out)
    return result
