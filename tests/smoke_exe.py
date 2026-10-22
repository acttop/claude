"""빌드된 exe 실제 실행 확인 (CI용). 실패 시 ::error:: 로 원인을 GitHub 화면에 표시한다."""
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

EXE = Path("dist/PDF2Excel.exe").resolve()
LOG = Path("smoke.log").resolve()
CASES = [("samples/sample.pdf", "smoke.xlsx", "P1_표1", "품목"),
         ("scan_sample.pdf", "smoke_scan.xlsx", "P1_표1", "번호")]


def fail(msg: str):
    log = LOG.read_text(encoding="utf-8", errors="replace") if LOG.exists() else "(로그 없음)"
    detail = f"{msg} | exe log: {log}".replace("\r", "").replace("\n", " / ")
    print(f"::error::{detail[:3500]}")
    sys.exit(1)


env = dict(os.environ, PDF2EXCEL_LOG=str(LOG))
for src, out, sheet, first in CASES:
    Path(out).unlink(missing_ok=True)
    p = subprocess.run([str(EXE), src, "-o", out], env=env, timeout=600)
    if not Path(out).exists():
        fail(f"{src}: 결과 파일 없음 (exit {p.returncode})")
    wb = load_workbook(out)
    if sheet not in wb.sheetnames:
        fail(f"{src}: 시트 {sheet} 없음, 실제 {wb.sheetnames}")
    rows = list(wb[sheet].iter_rows(values_only=True))
    if rows[0][0] != first and first not in rows[0]:
        fail(f"{src}: 헤더 불일치 {rows[:2]}")
    print(f"[OK] {src} -> {wb.sheetnames} {rows[:2]}")
print(LOG.read_text(encoding="utf-8") if LOG.exists() else "")
