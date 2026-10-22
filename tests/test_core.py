"""샘플 PDF를 생성해 변환 결과를 검증한다.  실행: python -m pytest tests"""
from pathlib import Path

import pytest
from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, PageBreak, TableStyle
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors

from pdf2excel.core import (LAYOUT_SINGLE, MODE_TEXT, ConvertOptions, convert,
                            parse_pages)

pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
KO = "HYSMyeongJo-Medium"

DATA = [["품목", "수량", "단가", "금액", "비율"],
        ["노트북", "10", "1,200,000", "12,000,000", "60%"],
        ["모니터", "20", "300,000", "6,000,000", "30%"],
        ["키보드", "40", "50,000", "2,000,000", "10%"]]


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    return make_sample_pdf(tmp_path / "sample.pdf")


def make_sample_pdf(path: Path) -> Path:
    """1쪽: 한글 표, 2쪽: 표 없는 텍스트."""
    style = ParagraphStyle("ko", fontName=KO, fontSize=11)
    table = Table(DATA)
    table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), KO),
                               ("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    doc = SimpleDocTemplate(str(path), pagesize=A4)
    doc.build([Paragraph("1쪽: 구매 내역", style), table, PageBreak(),
               Paragraph("2쪽: 표 없는 안내문입니다.", style),
               Paragraph("두 번째 줄 텍스트", style)])
    return path


def test_parse_pages():
    assert parse_pages("1-3,5") == [1, 2, 3, 5]
    assert parse_pages("") is None
    with pytest.raises(ValueError):
        parse_pages("0")


def test_auto_mode_per_page(sample_pdf, tmp_path):
    res = convert(sample_pdf, tmp_path / "out.xlsx")
    assert res.pages == 2 and res.tables == 1 and res.text_pages == 1
    wb = load_workbook(res.output)
    assert wb.sheetnames == ["P1_표1", "P2_텍스트"]
    ws = wb["P1_표1"]
    assert [c.value for c in ws[1]] == ["품목", "수량", "단가", "금액", "비율"]
    assert [c.value for c in ws[2]] == ["노트북", 10, 1200000, 12000000, 0.6]
    assert "안내문" in wb["P2_텍스트"]["A1"].value


def test_single_sheet_and_keep_text(sample_pdf, tmp_path):
    opts = ConvertOptions(layout=LAYOUT_SINGLE, convert_numbers=False)
    wb = load_workbook(convert(sample_pdf, tmp_path / "o.xlsx", opts).output)
    assert wb.sheetnames == ["전체"]
    values = [c.value for c in wb["전체"]["B"] if c.value]
    assert "1,200,000" not in values and "10" in values


def test_text_mode_page_select(sample_pdf, tmp_path):
    opts = ConvertOptions(mode=MODE_TEXT, pages=[1, 9])
    res = convert(sample_pdf, tmp_path / "t.xlsx", opts)
    assert res.pages == 1 and res.tables == 0
    assert any("존재하지 않는 페이지" in w for w in res.warnings)
