"""OCR 관련 테스트. OCR 엔진/학습 데이터가 없는 환경에서는 OCR 실행 테스트만 건너뛴다."""
from pathlib import Path

import pdfplumber
import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

from pdf2excel.analyze import METHOD_OCR, METHOD_OCR_FAILED, OCR_AUTO, analyze_page
from pdf2excel.core import convert
from pdf2excel.normalize import fix_amount, fix_code, fix_period, normalize_table

from make_scan import FONT, make_scan_pdf, sample_rows

try:
    from pdf2excel.extract_ocr import OcrEngine, find_tessdata
    find_tessdata()
    import tesserocr  # noqa: F401
    HAS_OCR = True
except Exception:  # noqa: BLE001
    HAS_OCR = False
needs_ocr = pytest.mark.skipif(not HAS_OCR, reason="OCR 엔진 또는 kor.traineddata 없음")


# ---------- normalize.py ----------
@pytest.mark.parametrize("raw, expected", [
    ("13,992:000", 13992000),   # 쉼표가 ':' 로 오인식
    ("1.200.000", 1200000),     # 쉼표가 '.' 로 오인식
    ("2,500;000", 2500000),
    ("13,992,000", 13992000),
    ("서버", "서버"),            # 숫자가 아니면 그대로
    ("12.5", "12.5"),            # 금액 열이 아니면 소수는 건드리지 않음
])
def test_fix_amount(raw, expected):
    assert fix_amount(raw) == expected


def test_fix_amount_in_amount_column():
    assert fix_amount("13 992:000", amount_column=True) == 13992000


@pytest.mark.parametrize("raw, expected", [
    ("202701-202912", "2027.01~2029.12"),
    ("2027.01-2029.12", "2027.01~2029.12"),
    ("2027.01 ~ 2029.12", "2027.01~2029.12"),
    ("202713-202912", "202713-202912"),   # 13월은 기간이 아님 → 그대로
    ("N-01", "N-01"),
])
def test_fix_period(raw, expected):
    assert fix_period(raw) == expected


def test_fix_code():
    assert fix_code("N-O7") == "N-07"
    assert fix_code("N-1l") == "N-11"


def test_normalize_table():
    rows = [["번호", "품목", "수량", "금액", "유지보수기간"],
            ["N-O1", "서버", "3", "13,992:000", "202701-202912"]]
    assert normalize_table(rows)[1] == ["N-01", "서버", "3", 13992000, "2027.01~2029.12"]


# ---------- analyze.py : 조용한 실패 금지 ----------
def _blank_scan(path: Path) -> Path:
    Image.new("L", (1240, 1754), 255).save(path, resolution=150)
    return path


def test_ocr_failed_when_engine_missing(tmp_path):
    pdf = make_scan_pdf(tmp_path / "s.pdf", [(sample_rows(1, 3), False)])
    with pdfplumber.open(pdf) as doc:
        r = analyze_page(doc.pages[0], 1, pdf_path=pdf, ocr=OCR_AUTO, engine=None,
                         engine_error="OCR 엔진(tesserocr)이 설치되어 있지 않습니다.")
    assert r.method == METHOD_OCR_FAILED
    assert "tesserocr" in r.error and "표가 0개" in r.error
    assert r.tables == []


@needs_ocr
def test_ocr_failed_on_page_without_table(tmp_path):
    pdf = _blank_scan(tmp_path / "blank.pdf")
    with pdfplumber.open(pdf) as doc, OcrEngine() as eng:
        r = analyze_page(doc.pages[0], 1, pdf_path=pdf, engine=eng)
    assert r.method == METHOD_OCR_FAILED and r.error


@needs_ocr
def test_convert_reports_ocr_failed(tmp_path):
    res = convert(_blank_scan(tmp_path / "blank.pdf"), tmp_path / "b.xlsx")
    assert res.failed_pages == [1] and res.methods[1] == METHOD_OCR_FAILED
    assert any("[ocr_failed]" in w for w in res.warnings)
    wb = load_workbook(res.output)
    assert wb.sheetnames == ["P1_OCR실패"]
    assert wb.active["A1"].value.startswith("[OCR 실패]")


# ---------- extract_ocr.py ----------
@needs_ocr
def test_scan_pdf_full(tmp_path):
    """방향 보정(2쪽 180도) + 표 밖 제목/쪽번호 제외 + 값 보정까지 원본과 일치해야 한다."""
    p1, p2 = sample_rows(1, 8), sample_rows(9, 6)
    pdf = make_scan_pdf(tmp_path / "scan.pdf", [(p1, False), (p2, True)])
    res = convert(pdf, tmp_path / "scan.xlsx")
    assert res.methods == {1: METHOD_OCR, 2: METHOD_OCR} and res.tables == 2
    assert any("180도" in w for w in res.warnings)

    wb = load_workbook(res.output)
    for sheet, rows in (("P1_표1", p1), ("P2_표1", p2)):
        got = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
        assert got[0] == ["번호", "품목", "수량", "금액", "유지보수기간"]
        expected = [[r[0], r[1], int(r[2]), int(r[3].replace(",", "")), r[4]] for r in rows]
        assert got[1:] == expected
        flat = " ".join(str(v) for r in got for v in r)
        assert "유지보수 목록" not in flat and "- 1 -" not in flat  # 표 밖 텍스트 제외


@needs_ocr
@pytest.mark.parametrize("size", [10, 18])
def test_single_digit_cell_reread(size):
    """넓은 셀 안의 작은 한 자리 수량: 빈 셀로 버리지 않고, 첫 인식 실패 시 2배 확대 재인식."""
    import numpy as np
    from pdf2excel.extract_ocr import DIGITS, _clean_cell, _clean_text, read_cell
    img = Image.new("L", (160, 90), 255)
    ImageDraw.Draw(img).text((70, 30), "3", font=ImageFont.truetype(str(FONT), size), fill=0)
    cell = np.array(img)
    with OcrEngine() as eng:
        if size == 10:  # 원래 크기로는 못 읽는 경우 → 재인식 경로 검증
            assert _clean_text(eng.read(_clean_cell(cell), 7, DIGITS)[0]) is None
        assert read_cell(eng, cell, kind="qty") == "3"


@needs_ocr
def test_speck_only_cell_is_empty():
    import numpy as np
    from pdf2excel.extract_ocr import read_cell
    rng = np.random.default_rng(0)
    cell = np.clip(255 + rng.integers(-25, 25, (110, 200)), 0, 255).astype(np.uint8)
    cell[rng.integers(0, 110, 15), rng.integers(0, 200, 15)] = 0  # 스캔 잡티
    with OcrEngine() as eng:
        assert read_cell(eng, cell, kind="qty") is None


# ---------- 2026-10-23 PM님 제보 사례 (워터마크 + 여러 줄 셀 + 한영 혼용) ----------
@needs_ocr
def test_maintenance_list_case(tmp_path):
    """제보 화면의 오류가 재발하지 않는지 확인한다.

    수정 전: "mH" 헤더, 구분 "0", 품번 "시-02"/"N-0S", 여러 줄 셀의 첫 줄 누락,
             "Gen9"→"6609", 워터마크 획 때문에 열이 늘어나고 행이 갈라짐
    """
    from difflib import SequenceMatcher
    from make_maint import HEADER, ROWS, make_maint_pdf

    res = convert(make_maint_pdf(tmp_path / "m.pdf"), tmp_path / "m.xlsx")
    got = [list(r) for r in load_workbook(res.output)["P1_표1"].iter_rows(values_only=True)]
    assert got[0] == HEADER
    assert len(got) == len(ROWS) + 1 and all(len(r) == 4 for r in got)  # 행/열 수 그대로
    for g, e in zip(got[1:], ROWS):
        assert g[0] == e[0] and g[1] == e[1] and g[3] == int(e[3])  # 품번·구분·수량 정확
        assert g[2].count("\n") == e[2].count("\n")                   # 여러 줄 유지
    must = ["EMC VNX 5200", "HP DL380 Gen9 8SFF CTO Server", "코어브릿지",
            "REDHAT LINUX E.E", "6식 라이선스 갱신 포함", "AXGATE 2300", "G-bic"]
    names = "\n".join(r[2] for r in got[1:])
    for m in must:
        assert m in names
    truth = "\n".join(r[2] for r in ROWS)
    assert SequenceMatcher(None, names, truth).ratio() >= 0.95


def test_code_column_fixes():
    """제보 화면의 품번 오인식: "시-02"(접두어), "N-0S"(S↔5), "0"(N-04 자리)."""
    from pdf2excel.normalize import fix_code_column
    got = fix_code_column(["N-01", "시-02", "N-03", "0", "N-0S"])
    assert got == ["N-01", "N-02", "N-03", "N-04", "N-05"]
    # 연속 번호 근거가 없으면 추측하지 않는다
    assert fix_code_column(["N-01", "0", "N-07"])[1] == "0"
